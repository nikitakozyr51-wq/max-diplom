#!/usr/bin/env python3
"""Build an editable university DOCX and render stable, paginated PDF."""
import argparse, json, os, re, shutil, subprocess, tempfile
from collections import Counter
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import fitz
from PIL import Image
from lxml import etree
ROOT=Path(__file__).resolve().parent
TEMPLATE=ROOT/'title-template.docx'
PLAN=ROOT/'plan.json'

def font(run,size=14,bold=None):
    run.font.name='Times New Roman';run.font.size=Pt(size)
    if bold is not None:run.bold=bold
    rpr=run._element.get_or_add_rPr();rf=rpr.rFonts
    if rf is not None:rf.set(qn('w:eastAsia'),'Times New Roman')

def replace(p,text,size=12):
    for r in list(p.runs):r._element.getparent().remove(r._element)
    font(p.add_run(text),size)

def build(data,out,pages):
    d=Document(TEMPLATE)
    plan=json.loads(PLAN.read_text(encoding='utf-8'))
    title=data.get('title',plan['title'])
    replace(d.paragraphs[0],'')
    replace(d.paragraphs[10],'Направление подготовки (специальность) 26.05.05 «Судовождение»')
    replace(d.paragraphs[23],'Группа С-612')
    replace(d.paragraphs[24],'Обучающегося Журилова Максима Максимовича')
    replace(d.paragraphs[26],'Вид работы: выпускная квалификационная работа специалиста')
    # Preserve two theme lines provided by the official source template.
    split=title.rfind(' ',0,72)
    replace(d.paragraphs[31],'Тема: '+title[:split],12)
    replace(d.paragraphs[32],title[split:].strip(),12)
    replace(d.paragraphs[45],'2026')
    # Keep typography of the official title; apply university body rules separately.
    s=d.add_section(WD_SECTION_START.NEW_PAGE)
    for sec in d.sections:
        sec.page_width=Mm(210);sec.page_height=Mm(297)
        sec.top_margin=sec.bottom_margin=Mm(20);sec.left_margin=Mm(30);sec.right_margin=Mm(15)
        sec.footer_distance=Mm(10)
    s.footer.is_linked_to_previous=False
    fp=s.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE')
    fr=fp.add_run('1');font(fr);fld.append(fr._r);fp._p.append(fld)
    body=d.styles.add_style('Thesis Body',1)
    body.font.name='Times New Roman';body.font.size=Pt(14)
    pf=body.paragraph_format;pf.line_spacing=1.5;pf.first_line_indent=Mm(12.5)
    pf.space_after=Pt(0);pf.space_before=Pt(0);pf.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    for n in (1,2,3):
        st=d.styles[f'Heading {n}'] if f'Heading {n}' in d.styles else d.styles.add_style(f'Heading {n}',1)
        ol=OxmlElement('w:outlineLvl');ol.set(qn('w:val'),str(n-1));st.element.get_or_add_pPr().append(ol)
        st.font.name='Times New Roman';st.font.size=Pt(14);st.font.bold=True;st.font.color.rgb=None;st.font.all_caps=(n==1)
        f=st.paragraph_format;f.line_spacing=1.5;f.space_before=Pt(12);f.space_after=Pt(12);f.keep_with_next=True;f.alignment=WD_ALIGN_PARAGRAPH.CENTER
        f.first_line_indent=Mm(0);f.page_break_before=n==1
    p=d.add_paragraph('ОГЛАВЛЕНИЕ');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(12)
    for r in p.runs:font(r,14,True)
    t=d.add_table(rows=0,cols=2);t.autofit=False;t.columns[0].width=Mm(151);t.columns[1].width=Mm(14)
    for item in plan['entries']:
        cells=t.add_row().cells;cells[0].width=Mm(151);cells[1].width=Mm(14)
        for j,txt in enumerate((item,str(pages[item]) if item in pages else '')):
            p=cells[j].paragraphs[0];p.paragraph_format.line_spacing=1.5;p.paragraph_format.space_after=Pt(0)
            p.paragraph_format.left_indent=Mm(5 if re.match(r'^\d\.\d',item) and j==0 else 0)
            p.alignment=WD_ALIGN_PARAGRAPH.RIGHT if j else WD_ALIGN_PARAGRAPH.LEFT
            font(p.add_run(txt))
        trpr=cells[0]._tc.getparent().get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
    for block in data['blocks']:
        typ=block['type']
        if typ=='heading':
            d.add_paragraph(block['text'],f"Heading {block.get('level',1)}")
        elif typ=='paragraph':
            p=d.add_paragraph(block['text'],'Thesis Body')
        elif typ=='equation':
            p=d.add_paragraph(style='Thesis Body');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Mm(0)
            if block.get('omml_xml'):
                math=etree.fromstring(block['omml_xml'].encode(),parser=etree.XMLParser(resolve_entities=False,no_network=True))
                if math.tag!=qn('m:oMath'):raise ValueError('omml_xml must have m:oMath root')
                p._p.append(math)
            elif block.get('omml',False):
                math=OxmlElement('m:oMath');mr=OxmlElement('m:r');mt=OxmlElement('m:t');mt.text=block['text'];mr.append(mt);math.append(mr);p._p.append(math)
            else:p.add_run(block['text'])
            if block.get('number'):p.add_run(f"    ({block['number']})")
        elif typ=='figure':
            with Image.open(block['path']) as im:iw,ih=im.size
            width=min(float(block.get('width_mm',165)),165,180*iw/ih)
            if width<=0:raise ValueError('Figure width must be positive')
            p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
            p.paragraph_format.space_after=Pt(0);p.add_run().add_picture(block['path'],width=Mm(width))
            p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1
            p.paragraph_format.space_after=Pt(6);p.paragraph_format.keep_together=True
            font(p.add_run(block['caption']),12)
        elif typ=='table':
            p=d.add_paragraph(block.get('caption',''),'Thesis Body');p.paragraph_format.first_line_indent=Mm(0);p.paragraph_format.keep_with_next=True
            headers=block.get('headers',[]);rows=([headers] if headers else [])+block['rows'];cols=len(rows[0])
            tb=d.add_table(rows=0,cols=cols);tb.autofit=False
            borders=OxmlElement('w:tblBorders')
            for edge in ('top','left','bottom','right','insideH','insideV'):
                line=OxmlElement('w:'+edge);line.set(qn('w:val'),'single');line.set(qn('w:sz'),'4');borders.append(line)
            tb._tbl.tblPr.append(borders)
            widths=block.get('widths_mm',[165/cols]*cols)
            if abs(sum(widths)-165)>1:raise ValueError('Table widths must sum to 165 mm')
            for col,w in zip(tb.columns,widths):col.width=Mm(w)
            for i,row in enumerate(rows):
                if len(row)!=cols:raise ValueError('Nonrectangular table')
                rr=tb.add_row()
                if not block.get('allow_row_split',False):rr._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                if i==0 and headers:
                    rr._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for j,(cell,txt) in enumerate(zip(rr.cells,row)):
                    cell.width=Mm(widths[j]);p=cell.paragraphs[0];p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                    p.paragraph_format.keep_with_next=i==0 and bool(headers)
                    font(p.add_run(str(txt)),12,i==0 and bool(headers))
        else:raise ValueError(f'Unknown block type: {typ}')
    if data.get('references'):
        d.add_paragraph('Список использованных источников','Heading 1')
        for i,ref in enumerate(data['references'],1):
            d.add_paragraph(f'{i} {ref}','Thesis Body')
    d.save(out)

def find_soffice():
    candidates = [os.environ.get('SOFFICE'), shutil.which('soffice'), shutil.which('libreoffice')]
    if os.name == 'nt':
        for base in ('PROGRAMFILES', 'PROGRAMFILES(X86)'):
            if os.environ.get(base):
                candidates.append(str(Path(os.environ[base])/'LibreOffice'/'program'/'soffice.exe'))
    candidates.append('/Applications/LibreOffice.app/Contents/MacOS/soffice')
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return str(Path(candidate).resolve())
    raise RuntimeError('LibreOffice not found. Install it or set SOFFICE to the soffice executable path.')

def render(docx,pdfdir):
    with tempfile.TemporaryDirectory(prefix='lo-thesis-') as profile:
        target = pdfdir/docx.with_suffix('.pdf').name
        if target.exists():
            target.unlink()
        result=subprocess.run([find_soffice(), '-env:UserInstallation='+Path(profile).resolve().as_uri(),
            '--headless','--convert-to','pdf:writer_pdf_Export','--outdir',str(pdfdir),str(docx)],
            capture_output=True,text=True,errors='replace',timeout=180)
        if result.returncode or not target.is_file():
            raise RuntimeError(result.stdout+result.stderr)

def rendered_font(pdf):
    fonts=Counter()
    with fitz.open(pdf) as document:
        for page in list(document)[3:]:
            for block in page.get_text('dict')['blocks']:
                for line in block.get('lines', []):
                    for span in line['spans']:
                        fonts[span['font']]+=len(span['text'])
    return fonts.most_common(1)[0][0] if fonts else 'unknown'

def paginate(pdf,headings):
    with fitz.open(pdf) as d:
        toc={title:page for level,title,page in d.get_toc()}
        missing=set(headings)-set(toc)
        if missing:raise RuntimeError('PDF bookmarks missing headings: '+repr(missing))
        return {h:toc[h] for h in headings},len(d)

def main():
    ap=argparse.ArgumentParser();ap.add_argument('input');ap.add_argument('output_stem');args=ap.parse_args()
    data=json.loads(Path(args.input).read_text(encoding='utf-8'));stem=Path(args.output_stem).resolve();stem.parent.mkdir(parents=True,exist_ok=True)
    for block in data['blocks']:
        if block['type']=='figure':block['path']=str((Path(args.input).resolve().parent/Path(block['path'])).resolve())
    docx=stem.with_suffix('.docx');pdf=stem.with_suffix('.pdf')
    headings=[b['text'] for b in data['blocks'] if b['type']=='heading']
    if data.get('references'):headings.append('Список использованных источников')
    if len(headings)!=len(set(headings)):raise ValueError('Duplicate headings')
    pages={}
    for iteration in range(1,5):
        build(data,docx,pages);render(docx,stem.parent);actual,count=paginate(pdf,headings)
        if actual==pages:break
        pages=actual
    else:raise RuntimeError('Pagination did not stabilize')
    match=rendered_font(pdf)
    report={'pages':count,'heading_pages':actual,'passes':iteration,'font_requested':'Times New Roman','font_resolved':match,'font_exact':'timesnewroman' in re.sub(r'[^a-z]', '', match.lower()),'title_template':str(TEMPLATE),'body_words':sum(len(b.get('text','').split()) for b in data['blocks'])}
    stem.with_suffix('.build-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(report,ensure_ascii=False,indent=2))
if __name__=='__main__':main()
